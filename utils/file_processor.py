import os
import re
import uuid
from docx import Document
import openpyxl
import pdfplumber
from io import BytesIO
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in {'docx', 'xlsx', 'pdf'}

def mask_cpf_cnpj(text, session_data):
    # Padrão para CPF: XXX.XXX.XXX-XX
    cpf_pattern = re.compile(r'\b(\d{3}\.\d{3}\.\d{3}-\d{2})\b')
    # Padrão para CNPJ: XX.XXX.XXX/XXXX-XX
    cnpj_pattern = re.compile(r'\b(\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2})\b')
    
    # Mascarar CPFs
    cpfs = cpf_pattern.findall(text)
    for cpf in cpfs:
        token = f"[CPF_{uuid.uuid4().hex[:8]}]"
        session_data[token] = cpf
        text = text.replace(cpf, token)
    
    # Mascarar CNPJs
    cnpjs = cnpj_pattern.findall(text)
    for cnpj in cnpjs:
        token = f"[CNPJ_{uuid.uuid4().hex[:8]}]"
        session_data[token] = cnpj
        text = text.replace(cnpj, token)
    
    return text

def mask_text(text, mask_words, session_data):
    # Mascarar palavras específicas
    for word in mask_words:
        pattern = re.compile(re.escape(word), re.IGNORECASE)
        matches = pattern.findall(text)
        for match in matches:
            token = f"[MASKED_{uuid.uuid4().hex[:8]}]"
            session_data[token] = match
            text = pattern.sub(token, text)
    
    # Detectar e mascarar e-mails
    email_pattern = re.compile(r'[\w\.-]+@[\w\.-]+\.\w+')
    emails = email_pattern.findall(text)
    for email in emails:
        token = f"[EMAIL_{uuid.uuid4().hex[:8]}]"
        session_data[token] = email
        text = text.replace(email, token)
    
    # Detectar e mascarar CPF e CNPJ
    text = mask_cpf_cnpj(text, session_data)
    
    return text

def restore_text(text, session_data):
    for token, original in session_data.items():
        text = text.replace(token, original)
    return text

def process_docx(file_path, mask_words, session_data, is_masking=True):
    doc = Document(file_path)
    for paragraph in doc.paragraphs:
        if is_masking:
            paragraph.text = mask_text(paragraph.text, mask_words, session_data)
        else:
            paragraph.text = restore_text(paragraph.text, session_data)
    
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

def process_xlsx(file_path, mask_words, session_data, is_masking=True):
    wb = openpyxl.load_workbook(file_path)
    for sheet in wb:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value:
                    if is_masking:
                        cell.value = mask_text(str(cell.value), mask_words, session_data)
                    else:
                        cell.value = restore_text(str(cell.value), session_data)
    
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer

def process_pdf(file_path, mask_words, session_data, is_masking=True):
    buffer = BytesIO()
    p = canvas.Canvas(buffer, pagesize=letter)
    
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                if is_masking:
                    masked_text = mask_text(text, mask_words, session_data)
                    p.drawString(100, 700, masked_text)
                else:
                    restored_text = restore_text(text, session_data)
                    p.drawString(100, 700, restored_text)
            p.showPage()
    
    p.save()
    buffer.seek(0)
    return buffer